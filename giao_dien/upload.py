"""
upload.py - Upload ảnh + file tài liệu.

Sửa: dai_nao.ghi_nho → luu_tru.ghi_nho.
"""

import io
import time
import secrets

from flask import session as phien_flask

from luu_tru.ghi_nho import (
    luu_file_gridfs,
    luu_metadata_anh_file,
    luu_noi_dung_trich_xuat,
    luu_lich_su_gui,
)


def _ghi_log(loai, noi_dung):
    try:
        from logs.ghi_log import ghi_log
        ghi_log(loai, noi_dung)
    except Exception:
        pass


def _lay_chu_so_huu():
    ten_tk = phien_flask.get("ten_dang_nhap")
    if ten_tk:
        return ten_tk
    return "khach"


def _tao_id():
    return "file-" + secrets.token_hex(8)


def _lay_duoi_file(ten_file):
    if not ten_file or "." not in ten_file:
        return ""
    return ten_file.rsplit(".", 1)[-1].lower()


def _trich_xuat_txt(noi_dung_bytes):
    try:
        return noi_dung_bytes.decode("utf-8", errors="replace")
    except Exception:
        return ""


def _trich_xuat_pdf(noi_dung_bytes):
    try:
        from PyPDF2 import PdfReader
        reader = PdfReader(io.BytesIO(noi_dung_bytes))
        ket_qua = []
        for trang in reader.pages[:50]:
            ket_qua.append(trang.extract_text() or "")
        return "\n".join(ket_qua).strip()
    except ImportError:
        return ""
    except Exception:
        return ""


def _trich_xuat_docx(noi_dung_bytes):
    try:
        from docx import Document
        doc = Document(io.BytesIO(noi_dung_bytes))
        return "\n".join(p.text for p in doc.paragraphs).strip()
    except ImportError:
        return ""
    except Exception:
        return ""


def _trich_xuat_xlsx(noi_dung_bytes):
    try:
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(noi_dung_bytes), read_only=True, data_only=True)
        ket_qua = []
        for sheet in wb.worksheets:
            ket_qua.append(f"[Sheet: {sheet.title}]")
            for hang in sheet.iter_rows(values_only=True):
                ket_qua.append("\t".join("" if c is None else str(c) for c in hang))
        return "\n".join(ket_qua).strip()
    except ImportError:
        return ""
    except Exception:
        return ""


def _trich_xuat_anh_gemini(noi_dung_bytes, duoi_file):
    try:
        import base64
        import requests
        from luu_tru.ghi_nho import lay_danh_sach_key_cua

        chu_so_huu = _lay_chu_so_huu()
        danh_sach = lay_danh_sach_key_cua(chu_so_huu) or []
        gemini_key = None
        for k in danh_sach:
            if k.get("provider") == "Gemini":
                gemini_key = k.get("key")
                break

        if not gemini_key:
            return ""

        duoi = duoi_file or "png"
        mime = "image/" + ("jpeg" if duoi == "jpg" else duoi)
        du_lieu_base64 = base64.b64encode(noi_dung_bytes).decode("utf-8")

        url = (
            "https://generativelanguage.googleapis.com/v1beta/models/"
            f"gemini-3.1-flash-lite:generateContent?key={gemini_key}"
        )
        body = {
            "contents": [{
                "parts": [
                    {"text": "Đọc và trích xuất toàn bộ nội dung trong ảnh này. "
                             "Nếu là code, giữ nguyên định dạng. Nếu là văn bản, "
                             "ghi lại đầy đủ. Chỉ trả về nội dung, không giải thích."},
                    {"inline_data": {"mime_type": mime, "data": du_lieu_base64}},
                ]
            }]
        }

        r = requests.post(url, json=body, timeout=30)
        if r.status_code != 200:
            return ""

        du_lieu = r.json()
        try:
            return du_lieu["candidates"][0]["content"]["parts"][0]["text"].strip()
        except (KeyError, IndexError):
            return ""
    except Exception:
        return ""


def _trich_xuat_noi_dung(noi_dung_bytes, duoi_file):
    if duoi_file == "txt":
        return _trich_xuat_txt(noi_dung_bytes)
    if duoi_file == "pdf":
        return _trich_xuat_pdf(noi_dung_bytes)
    if duoi_file in ("doc", "docx"):
        return _trich_xuat_docx(noi_dung_bytes)
    if duoi_file in ("xls", "xlsx", "csv"):
        return _trich_xuat_xlsx(noi_dung_bytes)
    if duoi_file in ("png", "jpg", "jpeg", "webp", "bmp", "gif"):
        return _trich_xuat_anh_gemini(noi_dung_bytes, duoi_file)
    return ""


def _xu_ly_mot_file(file_storage, loai_file, id_tro_chuyen=None, id_du_an=None):
    chu_so_huu = _lay_chu_so_huu()

    ten_file = file_storage.filename or "khong_ten"
    duoi_file = _lay_duoi_file(ten_file)

    try:
        file_storage.seek(0)
        noi_dung_bytes = file_storage.read()
    except Exception as e:
        return {"thanh_cong": False, "loi": f"Không đọc được file: {e}"}

    if not noi_dung_bytes:
        return {"thanh_cong": False, "loi": "File rỗng."}

    if len(noi_dung_bytes) > 20 * 1024 * 1024:
        return {"thanh_cong": False, "loi": "File quá lớn (tối đa 20MB)."}

    id_file = _tao_id()
    thoi_gian = int(time.time())

    meta_gridfs = {
        "id_file": id_file,
        "chu_so_huu": chu_so_huu,
        "loai_file": loai_file,
        "duoi_file": duoi_file,
        "thoi_gian": thoi_gian,
    }
    if id_tro_chuyen:
        meta_gridfs["id_tro_chuyen"] = id_tro_chuyen
    if id_du_an:
        meta_gridfs["id_du_an"] = id_du_an

    try:
        id_gridfs = luu_file_gridfs(
            ten_file=ten_file,
            noi_dung=noi_dung_bytes,
            metadata=meta_gridfs,
        )
    except Exception as e:
        return {"thanh_cong": False, "loi": f"Không lưu được file: {e}"}

    url = f"/api/file/{id_file}"

    metadata = {
        "id_file": id_file,
        "id_gridfs": str(id_gridfs),
        "ten_file": ten_file,
        "duoi_file": duoi_file,
        "kich_thuoc": len(noi_dung_bytes),
        "loai_file": loai_file,
        "chu_so_huu": chu_so_huu,
        "url": url,
        "thoi_gian": thoi_gian,
    }
    if id_tro_chuyen:
        metadata["id_tro_chuyen"] = id_tro_chuyen
    if id_du_an:
        metadata["id_du_an"] = id_du_an

    luu_metadata_anh_file(metadata)

    noi_dung_trich_xuat = _trich_xuat_noi_dung(noi_dung_bytes, duoi_file)
    if noi_dung_trich_xuat:
        du_lieu_trich_xuat = {
            "id_file": id_file,
            "chu_so_huu": chu_so_huu,
            "noi_dung": noi_dung_trich_xuat,
            "thoi_gian": thoi_gian,
        }
        if id_tro_chuyen:
            du_lieu_trich_xuat["id_tro_chuyen"] = id_tro_chuyen
        if id_du_an:
            du_lieu_trich_xuat["id_du_an"] = id_du_an
        luu_noi_dung_trich_xuat(du_lieu_trich_xuat)

    du_lieu_lich_su = {
        "id_file": id_file,
        "chu_so_huu": chu_so_huu,
        "ten_file": ten_file,
        "loai_file": loai_file,
        "thoi_gian": thoi_gian,
    }
    if id_tro_chuyen:
        du_lieu_lich_su["id_tro_chuyen"] = id_tro_chuyen
    if id_du_an:
        du_lieu_lich_su["id_du_an"] = id_du_an
    luu_lich_su_gui(du_lieu_lich_su)

    _ghi_log(
        loai_file,
        f"Upload {loai_file}: {ten_file} ({len(noi_dung_bytes)} bytes)"
    )

    return {
        "thanh_cong": True,
        "id": id_file,
        "url": url,
        "ten_file": ten_file,
        "co_noi_dung": bool(noi_dung_trich_xuat),
    }


def upload_anh(files, id_tro_chuyen=None, id_du_an=None):
    danh_sach_file = files.getlist("anh") if hasattr(files, "getlist") else []
    if not danh_sach_file:
        return {"thanh_cong": False, "loi": "Không có ảnh nào được gửi."}

    urls = []
    chi_tiet = []
    for f in danh_sach_file:
        ket_qua = _xu_ly_mot_file(f, "anh", id_tro_chuyen, id_du_an)
        if ket_qua.get("thanh_cong"):
            urls.append(ket_qua["url"])
        chi_tiet.append(ket_qua)

    return {
        "thanh_cong": len(urls) > 0,
        "urls": urls,
        "chi_tiet": chi_tiet,
    }


def upload_file(files, id_tro_chuyen=None, id_du_an=None):
    danh_sach_file = files.getlist("file") if hasattr(files, "getlist") else []
    if not danh_sach_file:
        return {"thanh_cong": False, "loi": "Không có file nào được gửi."}

    urls = []
    chi_tiet = []
    for f in danh_sach_file:
        ket_qua = _xu_ly_mot_file(f, "tai_lieu", id_tro_chuyen, id_du_an)
        if ket_qua.get("thanh_cong"):
            urls.append(ket_qua["url"])
        chi_tiet.append(ket_qua)

    return {
        "thanh_cong": len(urls) > 0,
        "urls": urls,
        "chi_tiet": chi_tiet,
    }